# -*- coding: utf-8 -*-
"""pptx_kit — 中性 PPT 技术函数库（python-pptx 之上的薄封装）。

本库**零视觉主张**：不内置任何色板、字体预设或页面骨架。
- 颜色一律由调用方显式传参；
- 字体通过 K.configure(cjk="…") 由项目声明（工具函数据此设置 a:ea 东亚字体与西文字体）；
- 页面骨架（眉题/标题/分隔线/页脚/金句带等）属于项目设计系统，请在项目脚本中自建
  （可参考 examples/minimal_deck.py 的做法与 reference.md 的历史项目示例）。

提供：页面尺寸常量、字体声明、文本、矩形/细线/虚线箭头/箭头形状/chip/占位框、
图片 cover 裁剪、生成后自检。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image, ImageOps

# ---------- 页面几何常量（事实值，非视觉主张） ----------
SW, SH = 13.333, 7.5      # 16:9
ML, MR = 0.7, 0.7         # 左右页边距
CW = SW - ML - MR         # 内容宽

# ---------- 项目字体声明 ----------
_FONTS = {"cjk": None}


def configure(cjk):
    """项目必须在生成前声明中文字体，如 K.configure(cjk="PingFang SC")。
    该字体将写入每个 run 的 a:latin / a:ea / a:cs（工具函数内部使用）。"""
    _FONTS["cjk"] = cjk


def _cjk():
    if not _FONTS["cjk"]:
        raise RuntimeError("请先调用 K.configure(cjk=\"…\") 声明项目字体")
    return _FONTS["cjk"]


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs, bg):
    """新建空白页。bg 必填（项目底色），传 None 表示不铺底。"""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if bg is not None:
        rect(s, 0, 0, SW, SH, fill=bg)
    return s


# ---------- 文本（字体/颜色由调用方提供） ----------
def set_run(run, size, color, bold=False, font=None, italic=False, spc=None):
    if font is None:
        font = _cjk()
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)
    if spc is not None:
        rPr.set("spc", str(spc))


def txbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, parts, size, color, bold=False, align=PP_ALIGN.LEFT,
         font=None, space_before=0, space_after=0, line=1.22, first=False, spc=None):
    """parts: str 或 [(text, color, bold[, size[, font]]), ...]；
    font 缺省用 configure() 声明的项目中文字体（单 run 内可再用元组第 5 位覆盖）。"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if isinstance(parts, str):
        parts = [(parts, color, bold)]
    for t in parts:
        text, c, b = t[0], t[1], t[2]
        sz = t[3] if len(t) > 3 else size
        fn = t[4] if len(t) > 4 else font
        r = p.add_run()
        r.text = text
        set_run(r, sz, c, b, fn, spc=spc)
    return p


# ---------- 形状（颜色一律显式传入；None = 无填充/无边框） ----------
def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, rounded=False, dash=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.055
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
        if dash:
            ln = shp.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    shp.text_frame.word_wrap = True
    return shp


def hairline(slide, x, y, w, color, weight=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    return c


def dashed(slide, x1, y1, x2, y2, color, weight=1.0, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if arrow:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    return c


def arrow_shape(slide, x, y, w, h, color, shape=MSO_SHAPE.RIGHT_ARROW):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def chip(slide, x, y, w, h, text, fill, tcolor, size=10.5, bold=True, line=None, font=None):
    shp = rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, color=tcolor, bold=bold, font=font,
         align=PP_ALIGN.CENTER, first=True, line=1.0)
    return shp


def placeholder(slide, x, y, w, h, label, fill, line_color, tcolor):
    """图片占位框（真实素材未就绪时用），颜色由调用方指定。"""
    rect(slide, x, y, w, h, fill=fill, line=line_color, lw=0.75, dash="dash")
    tf = txbox(slide, x + 0.15, y + h / 2 - 0.3, w - 0.3, 0.6, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, label, size=9.5, color=tcolor, align=PP_ALIGN.CENTER, first=True, line=1.2)


# ---------- 图片 ----------
def img_cover(src, dst, w_in, h_in, dpi=200, centering=(0.5, 0.5)):
    """中心裁剪(cover)到目标框比例，避免 add_picture 拉伸变形。"""
    if not os.path.exists(dst):
        im = Image.open(src).convert("RGB")
        im = ImageOps.fit(im, (int(w_in * dpi), int(h_in * dpi)),
                          Image.LANCZOS, centering=centering)
        im.save(dst, quality=92)
    return dst


def picture(slide, src, x, y, w, h, tmpdir="tmp_build", cover=True):
    if cover:
        os.makedirs(tmpdir, exist_ok=True)
        name = os.path.basename(src).rsplit(".", 1)[0] + f"_{int(w*100)}x{int(h*100)}.jpg"
        src = img_cover(src, os.path.join(tmpdir, name), w, h)
    return slide.shapes.add_picture(src, Inches(x), Inches(y), Inches(w), Inches(h))


# ---------- 自检 ----------
def selfcheck(path, banned=(), extra_assert=None):
    """生成后自检：禁用词断言。返回文本全集供进一步断言。"""
    chk = Presentation(path)
    blob = "\n".join(sh.text_frame.text
                     for sl in chk.slides for sh in sl.shapes if sh.has_text_frame)
    for w in banned:
        assert w not in blob, f"禁用词命中: {w}"
    if extra_assert:
        extra_assert(blob)
    n = len(chk.slides._sldIdLst)
    print(f"selfcheck OK — slides: {n}")
    return blob
